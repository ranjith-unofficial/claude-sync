#!/usr/bin/env python3
"""Turn the two Inc42 tracking spreadsheets into machine-readable JSON assertions.

Inputs (override with --target-spec / --master):
  "Inc42 Website - Events & Properties.xlsx"  -> the TARGET taxonomy (snake_case, v1.0, aspirational)
  "Inc42 Analytics _ Master.xlsx"             -> the LIVE taxonomy (Title Case, what is actually deployed)

Outputs:
  spec/target.json    58 spec'd events + property groups + dictionary (what SHOULD exist)
  spec/live.json      observed live events w/ 30d volume + known findings (what DOES exist)
  spec/aliases.json   target event <-> live event name map (hand-maintained after first generation)
"""
import argparse, json, os, re, sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("pip3 install openpyxl")

HOME = Path.home()
DEF_TARGET = HOME / "Downloads" / "Inc42 Website - Events & Properties.xlsx"
DEF_MASTER = HOME / "Downloads" / "Inc42 Analytics _ Master.xlsx"
OUT = Path(__file__).resolve().parent.parent / "spec"

def cells(row):
    return [("" if c is None else str(c).replace("\n", " ").strip()) for c in row]

def split_enum(text):
    """'a (x/y/z) · b' -> [{'name':'a','values':['x','y','z']}, {'name':'b','values':[]}]"""
    out = []
    for part in re.split(r"\s*[·|]\s*", text or ""):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^([A-Za-z0-9_\[\]\.\/ ]+?)\s*\((.*)\)\s*$", part)
        if m:
            name = m.group(1).strip()
            vals = [v.strip() for v in re.split(r"\s*/\s*", m.group(2)) if v.strip()]
        else:
            name, vals = part, []
        name = name.replace("[]", "").strip()
        if name:
            out.append({"name": name, "values": vals})
    return out

def parse_target(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Website | Events"]
    rows = [cells(r) for r in ws.iter_rows(values_only=True)]
    header = rows[0]
    idx = {h.strip(): i for i, h in enumerate(header) if h.strip()}

    def col(r, key, default=""):
        i = idx.get(key)
        return r[i] if i is not None and i < len(r) else default

    events, group = [], None
    for r in rows[1:]:
        if not any(r):
            continue
        name = col(r, "Event").strip()
        grp = col(r, "Group").strip()
        # group header rows carry a Group but no Event
        if grp and not name:
            group = grp
            continue
        if not name or name.startswith("★"):
            continue
        group = grp or group
        prop_groups = [g.strip() for g in re.split(r"\s*\+\s*", col(r, "Property Groups")) if g.strip()]
        dests = {d: bool(col(r, d).strip()) for d in ("PostHog", "Customer.io", "GA4", "Meta Pixel")}
        events.append({
            "event": name,
            "group": group,
            "fires_when": col(r, "Fires when"),
            "property_groups": [re.sub(r"\*$", "", g) for g in prop_groups],
            "story_conditional": "*" in col(r, "Property Groups"),
            "properties": split_enum(col(r, "Event properties (enums inline)")),
            "person_properties": [p["name"] for p in split_enum(col(r, "Person Property to Update"))],
            "calls_identify": bool(col(r, "Identify").strip()),
            "destinations": dests,
            "metric": col(r, "Metric / trigger served"),
            "notes": col(r, "Notes"),
            "critical": "★" in col(r, "Event") or "★" in col(r, "Metric / trigger served"),
            "upsell": "⬆" in col(r, "Metric / trigger served"),
        })

    # property architecture
    pa = [cells(r) for r in wb["Website | Property Architecture"].iter_rows(values_only=True)]
    groups, dictionary, mode = {}, [], "groups"
    for r in pa:
        if not any(r):
            continue
        if r[0].startswith("PROPERTY DICTIONARY"):
            mode = "dict"; continue
        if r[0] in ("Group", "Property"):
            continue
        if mode == "groups" and r[0] in ("Super", "Story", "Person"):
            groups[r[0]] = {
                "properties": [p["name"] for p in split_enum(r[1])],
                "attached_to": r[2] if len(r) > 2 else "",
            }
        elif mode == "dict" and r[0]:
            names = [n.strip() for n in re.split(r"\s*/\s*", r[0]) if n.strip()]
            allowed = r[3] if len(r) > 3 else ""
            vals = [] if ("CONTEXT-DEPENDENT" in allowed or "…" in allowed and "/" not in allowed) else \
                   [v.strip() for v in re.split(r"\s*/\s*", re.sub(r"\s*\(.*?\)", "", allowed)) if v.strip() and len(v.strip()) < 40]
            for n in names:
                dictionary.append({
                    "property": n,
                    "group": r[1] if len(r) > 1 else "",
                    "type": r[2] if len(r) > 2 else "",
                    "allowed": vals,
                    "raw_allowed": allowed,
                })
    return {"source": os.path.basename(path), "version": "web-spec-v1.0",
            "events": events, "property_groups": groups, "dictionary": dictionary}

def parse_master(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    live, enums, findings = [], [], []
    ws = wb["Audit | Jun 2026"]
    rows = [cells(r) for r in ws.iter_rows(values_only=True)]
    mode = None
    for r in rows:
        if not any(r):
            continue
        head = r[0]
        if head.startswith("EVENT-LEVEL"):
            mode = "events"; continue
        if head.startswith("ENUM VALUE"):
            mode = "enums"; continue
        if head in ("Event (Plan)", "Property"):
            continue
        if mode == "events" and len(r) >= 4:
            live_name = r[1].strip()
            if not live_name or live_name in ("(absent)", "(property only)"):
                live_name = None
            try:
                vol = float(r[2]) if r[2] else 0.0
            except ValueError:
                vol = 0.0
            status = r[3]
            live.append({
                "plan_name": r[0].strip(),
                "live_name": live_name,
                "volume_30d": vol,
                "status": status,
                "healthy": "✅" in status,
                "dead": "Dead" in status or "Near-zero" in status,
                "missing": "Missing" in status or live_name is None,
                "duplicate": "dup" in status.lower(),
                "noise": "Noise" in status,
                "name_drift": "drift" in status.lower(),
                "findings": r[4] if len(r) > 4 else "",
                "priority": r[5] if len(r) > 5 else "",
                "fix": r[6] if len(r) > 6 else "",
            })
        elif mode == "enums" and len(r) >= 2:
            enums.append({
                "property": r[0].strip(),
                "plan_values_present": r[1] if len(r) > 1 else "",
                "drift": r[2] if len(r) > 2 else "",
                "junk_values": [v.strip() for v in re.split(r",\s*", r[3]) if v.strip() and v.strip() != "—"] if len(r) > 3 else [],
            })

    props = []
    pw = [cells(r) for r in wb["Property Audit | Jun 2026"].iter_rows(values_only=True)]
    for r in pw:
        if not any(r) or r[0] in ("Group", "") or r[0].startswith("Inc42 PostHog") or r[0].startswith("UNDOCUMENTED"):
            continue
        if len(r) < 6:
            continue
        props.append({
            "group": r[0], "plan_property": r[1], "live_keys": r[2],
            "scope": r[3], "coverage": r[4], "status": r[5],
            "notes": r[6] if len(r) > 6 else "",
            "ok": "✅" in r[5], "warn": "⚠️" in r[5], "dead": "✗" in r[5],
        })
    return {"source": os.path.basename(path), "as_of": "2026-06-22", "project_id": 53557,
            "events": live, "enum_conformance": enums, "properties": props}

def build_aliases(target, live):
    """Seed the target<->live name map. Exact + normalised match, rest left null for a human."""
    def norm(s):
        return re.sub(r"[^a-z0-9]", "", (s or "").lower())
    live_by_norm = {norm(e["live_name"]): e["live_name"] for e in live["events"] if e["live_name"]}
    aliases = {}
    for ev in target["events"]:
        aliases[ev["event"]] = {"live": live_by_norm.get(norm(ev["event"])), "confidence": "auto" if live_by_norm.get(norm(ev["event"])) else "unmapped"}
    return aliases

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--target-spec", default=str(DEF_TARGET))
    ap.add_argument("--master", default=str(DEF_MASTER))
    a = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)

    target = parse_target(a.target_spec)
    live = parse_master(a.master)
    (OUT / "target.json").write_text(json.dumps(target, indent=2, ensure_ascii=False))
    (OUT / "live.json").write_text(json.dumps(live, indent=2, ensure_ascii=False))
    ap_path = OUT / "aliases.json"
    if not ap_path.exists():
        ap_path.write_text(json.dumps(build_aliases(target, live), indent=2, ensure_ascii=False))
        seeded = "seeded"
    else:
        seeded = "kept existing"
    print(f"target.json  {len(target['events'])} events, {len(target['dictionary'])} dictionary entries")
    print(f"live.json    {len(live['events'])} live events, {len(live['properties'])} properties, {len(live['enum_conformance'])} enum rows")
    print(f"aliases.json {seeded}")
